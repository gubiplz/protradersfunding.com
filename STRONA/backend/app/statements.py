"""Wyciąg z kont — jeden skoroszyt Excela z całą historią, gotowy dla AI.

Broker daje klientowi „statement": listę transakcji, krzywą kapitału i podsumowanie
w jednym pliku. Tutaj to samo, tyle że za jednym kliknięciem dla WSZYSTKICH kont
naraz — bo pytanie, które padnie nad tym plikiem, brzmi zwykle „jak nam idzie",
a nie „jak idzie temu jednemu".

Układ jest podporządkowany temu, że czytelnikiem bywa model językowy:

* **Płaskie, długie tabele, nie arkusz per konto.** Arkusz per konto wygląda
  jak wyciąg z banku i jest nie do przetworzenia: żeby policzyć cokolwiek na
  przekroju, trzeba najpierw skleić trzydzieści arkuszy. Każdy wiersz niesie
  więc `account_login`, a konto jest KOLUMNĄ, nie zakładką.
* **Arkusz `Dictionary` na pierwszym miejscu.** Model, który dostaje `pnl`
  i `overall_dd_used_pct` bez wyjaśnienia, zgaduje jednostki. Słownik mówi
  wprost, co znaczy każda kolumna i w czym jest liczona.
* **Liczby jako liczby, daty jako daty.** Kwota zapisana tekstem („$1,234.00")
  wymaga parsowania i potrafi zostać przeczytana jako 1.234.
* **Nic nie jest liczone dwa razy inaczej.** Metryki reguł pochodzą z
  `rules.display_metrics` — z tej samej funkcji, która karmi panel i portal.

Konta DARMOWE są poza zakresem: nie prowadzimy ich tak jak płatnych, więc ich
historia nie opowiada nic o tym, jak idzie usługa. Ten sam filtr co domyślna
lista w panelu — patrz `konta_do_wyciagu`.
"""
from __future__ import annotations

import io
from collections import defaultdict
from datetime import datetime, timezone

from . import rules
from .models import Account, Breach, EquitySnapshot, Payout, Trade, Trader

# Nazwa arkusza w Excelu: max 31 znaków, bez : \\ / ? * [ ]
ARKUSZE = ("Dictionary", "Accounts", "Trades", "Daily", "Symbols",
           "Equity", "Payouts", "Breaches")

PIENIADZ = '#,##0.00'
PROCENT = '0.00'
DATA = 'yyyy-mm-dd hh:mm'
DZIEN = 'yyyy-mm-dd'


# --------------------------------------------------------------------------- #
#  Wybór kont                                                                  #
# --------------------------------------------------------------------------- #
def konta_do_wyciagu(session, *, placacy_id, filtr_nie_import):
    """Konta, które trafiają do wyciągu: płacący klienci, bez importów.

    `placacy_id` i `filtr_nie_import` wstrzykuje `main`, bo to tam mieszkają —
    ten moduł nie ma prawa mieć DRUGIEJ definicji tego, kto zapłacił.

    Pula (trader_id NULL) odpada świadomie, w odróżnieniu od listy w panelu:
    tam widać ją po to, żeby kolejka provisioningu nie znikała z oczu, a tutaj
    byłby to wiersz o rachunku, za którym nie stoi żaden człowiek.
    """
    return (session.query(Account)
            .filter(Account.trader_id.isnot(None),
                    Account.trader_id.in_(placacy_id),
                    filtr_nie_import)
            .order_by(Account.created_at.desc(), Account.id.desc())
            .all())


# --------------------------------------------------------------------------- #
#  Liczenie                                                                    #
# --------------------------------------------------------------------------- #
def _naiwna(dt):
    """Baza trzyma czas bez strefy. Excel też nie chce strefy — a mieszanie
    jednego z drugim rzuca `TypeError` przy pierwszym porównaniu."""
    if dt is None:
        return None
    return dt.replace(tzinfo=None) if dt.tzinfo else dt


def statystyki_trejdow(trejdy: list[Trade]) -> dict:
    """Podsumowanie, które broker drukuje na pierwszej stronie wyciągu.

    Liczone WYŁĄCZNIE z pozycji zamkniętych: otwarta ma wynik pływający, więc
    wliczona do win rate policzyłaby jako wygraną coś, co jeszcze nie wygrało.
    """
    zamkniete = [t for t in trejdy if t.status == "closed"]
    wygrane = [t for t in zamkniete if (t.pnl or 0) > 0]
    przegrane = [t for t in zamkniete if (t.pnl or 0) < 0]
    zysk = sum(t.pnl or 0 for t in wygrane)
    strata = sum(t.pnl or 0 for t in przegrane)

    # Najdłuższe serie — jedyna liczba w tym zestawie, której nie da się
    # odtworzyć z sum, a mówi o zmienności więcej niż średnia.
    seria_w = seria_p = max_w = max_p = 0
    for t in sorted(zamkniete, key=lambda x: _naiwna(x.closed_at) or datetime.min):
        if (t.pnl or 0) > 0:
            seria_w, seria_p = seria_w + 1, 0
            max_w = max(max_w, seria_w)
        elif (t.pnl or 0) < 0:
            seria_p, seria_w = seria_p + 1, 0
            max_p = max(max_p, seria_p)

    otwarte = [t for t in trejdy if t.status != "closed"]
    return {
        "trades_total": len(trejdy),
        "trades_closed": len(zamkniete),
        "trades_open": len(otwarte),
        "wins": len(wygrane),
        "losses": len(przegrane),
        "win_rate_pct": round(len(wygrane) / len(zamkniete) * 100, 2) if zamkniete else None,
        "gross_profit": round(zysk, 2),
        "gross_loss": round(strata, 2),
        "net_pnl": round(zysk + strata, 2),
        # Brak strat to nie „nieskończony" profit factor, tylko brak mianownika.
        "profit_factor": round(zysk / abs(strata), 2) if strata else None,
        "avg_win": round(zysk / len(wygrane), 2) if wygrane else None,
        "avg_loss": round(strata / len(przegrane), 2) if przegrane else None,
        "largest_win": round(max((t.pnl or 0) for t in wygrane), 2) if wygrane else None,
        "largest_loss": round(min((t.pnl or 0) for t in przegrane), 2) if przegrane else None,
        "expectancy": round((zysk + strata) / len(zamkniete), 2) if zamkniete else None,
        "max_win_streak": max_w,
        "max_loss_streak": max_p,
        "total_lots": round(sum(t.lots or 0 for t in trejdy), 2),
        "first_trade_at": _naiwna(min((t.opened_at for t in trejdy), default=None)),
        "last_trade_at": _naiwna(max((_naiwna(t.closed_at) or _naiwna(t.opened_at)
                                      for t in trejdy), default=None)),
    }


def obsuniecie(krzywa: list[EquitySnapshot]) -> dict:
    """Największe obsunięcie kapitału — od szczytu do dołka, w kolejności czasu.

    To NIE jest `overall_dd_used_pct` z reguł: tamto mierzy odległość od podłogi,
    przy której konto pęka, i liczy się względem salda startowego. To tutaj jest
    historyczne, liczone od szczytu — dwie różne liczby o podobnej nazwie.
    """
    szczyt = None
    max_spadek = 0.0
    max_spadek_pct = 0.0
    kiedy = None
    for s in sorted(krzywa, key=lambda x: _naiwna(x.ts) or datetime.min):
        e = s.equity or 0
        if szczyt is None or e > szczyt:
            szczyt = e
            continue
        spadek = szczyt - e
        if spadek > max_spadek:
            max_spadek = spadek
            max_spadek_pct = spadek / szczyt * 100 if szczyt else 0
            kiedy = _naiwna(s.ts)
    return {"max_drawdown": round(max_spadek, 2),
            "max_drawdown_pct": round(max_spadek_pct, 2),
            "max_drawdown_at": kiedy,
            "peak_equity": round(szczyt, 2) if szczyt is not None else None}


def dni_handlowe(trejdy: list[Trade], krzywa: list[EquitySnapshot]) -> list[dict]:
    """Jeden wiersz na dzień i konto — to, co broker nazywa „daily summary".

    Dzień bierzemy z `day_key` snapshotu, a nie z daty kalendarzowej UTC:
    doba handlowa zaczyna się o godzinie serwera i to według niej liczą się
    limity dzienne. Dwie różne doby dałyby wynik dnia niezgodny z tym, na
    podstawie czego konto mogło polec.
    """
    po_dniach: dict[str, dict] = defaultdict(
        lambda: {"trades": 0, "wins": 0, "losses": 0, "gross_profit": 0.0,
                 "gross_loss": 0.0, "lots": 0.0, "symbols": set()})

    dzien_snapshotu: dict[str, list[EquitySnapshot]] = defaultdict(list)
    for s in krzywa:
        klucz = s.day_key or (_naiwna(s.ts).strftime("%Y-%m-%d") if s.ts else "")
        if klucz:
            dzien_snapshotu[klucz].append(s)

    # Trejd należy do dnia, w którym został ZAMKNIĘTY: wtedy wynik stał się
    # faktem. Pozycja otwarta nie ma jeszcze swojego dnia.
    granice = sorted((k, min(_naiwna(s.ts) for s in v)) for k, v in dzien_snapshotu.items())
    for t in trejdy:
        if t.status != "closed" or not t.closed_at:
            continue
        kiedy = _naiwna(t.closed_at)
        klucz = ""
        for k, poczatek in granice:
            if poczatek <= kiedy:
                klucz = k
            else:
                break
        klucz = klucz or kiedy.strftime("%Y-%m-%d")
        d = po_dniach[klucz]
        d["trades"] += 1
        d["lots"] += t.lots or 0
        d["symbols"].add(t.symbol)
        if (t.pnl or 0) > 0:
            d["wins"] += 1
            d["gross_profit"] += t.pnl or 0
        elif (t.pnl or 0) < 0:
            d["losses"] += 1
            d["gross_loss"] += t.pnl or 0

    out = []
    szczyt = None
    for klucz in sorted(set(list(po_dniach.keys()) + list(dzien_snapshotu.keys()))):
        snapy = sorted(dzien_snapshotu.get(klucz, []),
                       key=lambda x: _naiwna(x.ts) or datetime.min)
        d = po_dniach.get(klucz, {"trades": 0, "wins": 0, "losses": 0,
                                  "gross_profit": 0.0, "gross_loss": 0.0,
                                  "lots": 0.0, "symbols": set()})
        otwarcie = snapy[0].equity if snapy else None
        zamkniecie = snapy[-1].equity if snapy else None
        if zamkniecie is not None:
            szczyt = zamkniecie if szczyt is None else max(szczyt, zamkniecie)
        out.append({
            "day": klucz,
            "trades": d["trades"], "wins": d["wins"], "losses": d["losses"],
            "gross_profit": round(d["gross_profit"], 2),
            "gross_loss": round(d["gross_loss"], 2),
            "net_pnl": round(d["gross_profit"] + d["gross_loss"], 2),
            "lots": round(d["lots"], 2),
            "symbols": ", ".join(sorted(d["symbols"])),
            "equity_open": round(otwarcie, 2) if otwarcie is not None else None,
            "equity_close": round(zamkniecie, 2) if zamkniecie is not None else None,
            "balance_close": round(snapy[-1].balance, 2) if snapy else None,
            "day_return_pct": (round((zamkniecie - otwarcie) / otwarcie * 100, 3)
                               if otwarcie else None),
            "peak_equity_to_date": round(szczyt, 2) if szczyt is not None else None,
            "drawdown_from_peak_pct": (round((szczyt - zamkniecie) / szczyt * 100, 3)
                                       if szczyt and zamkniecie is not None else None),
            "snapshots": len(snapy),
        })
    return out


# --------------------------------------------------------------------------- #
#  Skoroszyt                                                                   #
# --------------------------------------------------------------------------- #
SLOWNIK = [
    ("Accounts", "account_login", "Login rachunku na platformie — klucz łączący wszystkie arkusze", "tekst"),
    ("Accounts", "initial_balance", "Saldo startowe rachunku", "USD"),
    ("Accounts", "balance / equity", "Saldo zamknięte / kapitał z wynikiem pozycji otwartych", "USD"),
    ("Accounts", "profit_pct", "Zysk względem salda startowego", "%"),
    ("Accounts", "max_drawdown_pct", "Największe historyczne obsunięcie od szczytu kapitału", "%"),
    ("Accounts", "overall_dd_used_pct", "Ile wykorzystano dopuszczalnego obsunięcia REGULAMINOWEGO (100 = rachunek pęka)", "%"),
    ("Accounts", "daily_loss_used_pct", "To samo dla limitu dziennego", "%"),
    ("Accounts", "profit_factor", "Suma zysków / wartość bezwzględna sumy strat; puste = brak strat", "krotność"),
    ("Accounts", "expectancy", "Średni wynik pojedynczej zamkniętej pozycji", "USD"),
    ("Accounts", "trading_days", "Dni z aktywnością, liczone regułami wyzwania", "dni"),
    ("Trades", "side", "buy = pozycja długa, sell = krótka", "tekst"),
    ("Trades", "lots", "Wolumen w lotach", "loty"),
    ("Trades", "pnl", "Wynik pozycji; dla otwartej jest to wynik pływający", "USD"),
    ("Trades", "duration_min", "Czas trzymania pozycji", "minuty"),
    ("Trades", "status", "closed = rozliczona, open = wciąż otwarta", "tekst"),
    ("Daily", "day", "Doba handlowa wg czasu serwera, nie kalendarzowa UTC", "data"),
    ("Daily", "net_pnl", "Wynik dnia z pozycji zamkniętych tego dnia", "USD"),
    ("Daily", "drawdown_from_peak_pct", "Odległość od najwyższego kapitału do tego dnia włącznie", "%"),
    ("Equity", "ts", "Znacznik czasu odczytu kapitału (UTC)", "data i godzina"),
    ("Equity", "open_pnl", "Wynik pływający pozycji otwartych w chwili odczytu", "USD"),
    ("Symbols", "net_pnl", "Wynik na instrumencie, pozycje zamknięte", "USD"),
    ("Payouts", "profit_amount", "Zysk wypłacany z rachunku, PRZED podziałem", "USD"),
    ("Payouts", "trader_share", "Część, która trafiła do tradera po podziale zysku", "USD"),
    ("Payouts", "paid", "PRAWDA = wypłata rozliczona", "tak/nie"),
    ("Accounts", "payouts_total", "Suma trader_share ze wszystkich wypłat rachunku", "USD"),
    ("Breaches", "type", "Złamana reguła: daily_loss, max_drawdown, consistency, manual", "tekst"),
]

NOTA = [
    "Wyciąg obejmuje rachunki płacących klientów. Darmowe rejestracje i rachunki",
    "importowane są POMINIĘTE — nie prowadzimy ich tak jak płatnych, więc ich",
    "historia nie mówi nic o tym, jak idzie usługa.",
    "",
    "Każdy wiersz w każdym arkuszu niesie kolumnę account_login. Konto jest",
    "kolumną, nie zakładką — dzięki temu da się liczyć na przekroju wszystkich",
    "rachunków bez sklejania arkuszy.",
    "",
    "Uwaga na dwie podobnie brzmiące liczby: max_drawdown_pct to historyczne",
    "obsunięcie od szczytu kapitału, a overall_dd_used_pct mówi, ile zużyto",
    "limitu regulaminowego, przy którym rachunek przestaje istnieć.",
]


def _naglowek(ws, kolumny, szerokosci=None):
    from openpyxl.styles import Alignment, Font, PatternFill
    ws.append(list(kolumny))
    for i, nazwa in enumerate(kolumny, start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F3B52")
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[c.column_letter].width = (
            (szerokosci or {}).get(nazwa, max(11, min(26, len(str(nazwa)) + 4))))
    ws.freeze_panes = "A2"


def _wiersze(ws, wiersze, formaty: dict[str, str]):
    """Dopisuje wiersze i nadaje format kolumnom po NAZWIE nagłówka."""
    naglowki = [c.value for c in ws[1]]
    for w in wiersze:
        ws.append([w.get(k) for k in naglowki])
    for i, nazwa in enumerate(naglowki, start=1):
        fmt = formaty.get(nazwa)
        if not fmt:
            continue
        for wiersz in range(2, ws.max_row + 1):
            ws.cell(row=wiersz, column=i).number_format = fmt
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def zbuduj(session, konta: list[Account]) -> bytes:
    """Cały skoroszyt w pamięci. Zwraca bajty gotowe do odesłania."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    ids = [a.id for a in konta]
    maile = dict(session.query(Trader.id, Trader.email).all())

    trejdy_konta: dict[int, list[Trade]] = defaultdict(list)
    if ids:
        for t in (session.query(Trade).filter(Trade.account_id.in_(ids))
                  .order_by(Trade.opened_at).all()):
            trejdy_konta[t.account_id].append(t)

    krzywe: dict[int, list[EquitySnapshot]] = defaultdict(list)
    if ids:
        for s in (session.query(EquitySnapshot)
                  .filter(EquitySnapshot.account_id.in_(ids))
                  .order_by(EquitySnapshot.ts).all()):
            krzywe[s.account_id].append(s)

    wyplaty: dict[int, list[Payout]] = defaultdict(list)
    breache: dict[int, list[Breach]] = defaultdict(list)
    if ids:
        for p in session.query(Payout).filter(Payout.account_id.in_(ids)).all():
            wyplaty[p.account_id].append(p)
        for b in session.query(Breach).filter(Breach.account_id.in_(ids)).all():
            breache[b.account_id].append(b)

    wb = Workbook()
    wb.remove(wb.active)

    # --- Dictionary ---------------------------------------------------------
    ws = wb.create_sheet("Dictionary")
    _naglowek(ws, ["sheet", "column", "meaning", "unit"],
              {"meaning": 78, "column": 24, "unit": 16})
    for arkusz, kolumna, opis, jednostka in SLOWNIK:
        ws.append([arkusz, kolumna, opis, jednostka])
    ws.append([])
    for linia in NOTA:
        ws.append([linia])
        ws.cell(row=ws.max_row, column=1).font = Font(italic=True)

    # --- Accounts -----------------------------------------------------------
    ws = wb.create_sheet("Accounts")
    _naglowek(ws, [
        "account_login", "trader_name", "trader_email", "product_key", "phase",
        "status", "source", "initial_balance", "balance", "equity", "open_pnl",
        "profit_pct", "max_drawdown", "max_drawdown_pct", "max_drawdown_at",
        "peak_equity", "daily_loss_used_pct", "overall_dd_used_pct",
        "target_equity", "daily_floor", "overall_floor", "profit_target_pct",
        "max_daily_loss_pct", "max_overall_loss_pct", "drawdown_type",
        "min_trading_days", "trading_days", "trades_total", "trades_closed",
        "trades_open", "wins", "losses", "win_rate_pct", "gross_profit",
        "gross_loss", "net_pnl", "profit_factor", "expectancy", "avg_win",
        "avg_loss", "largest_win", "largest_loss", "max_win_streak",
        "max_loss_streak", "total_lots", "first_trade_at", "last_trade_at",
        "payouts_count", "payouts_total", "breaches_count", "breach_reason",
        "bot_enabled", "mt5_backed", "platform_server", "created_at", "started_at",
    ], {"trader_email": 30, "trader_name": 22, "breach_reason": 34,
        "max_drawdown_at": 18, "first_trade_at": 18, "last_trade_at": 18,
        "created_at": 18, "started_at": 18})

    wiersze_kont = []
    for a in konta:
        trejdy = trejdy_konta.get(a.id, [])
        krzywa = krzywe.get(a.id, [])
        st = statystyki_trejdow(trejdy)
        dd = obsuniecie(krzywa)
        try:
            cfg = rules.config_from_account(a)
            m = rules.display_metrics(
                cfg, balance=a.balance or 0, equity=a.equity or 0,
                peak_equity=a.peak_equity or a.initial_balance or 0,
                day_start_equity=a.day_start_equity or a.initial_balance or 0,
                trading_days=a.trading_days_count or 0)
        except Exception:
            # Wyciąg nie może paść przez jedno konto o nietypowej konfiguracji.
            m = {}
        wyp = wyplaty.get(a.id, [])
        wiersze_kont.append({
            "account_login": a.login, "trader_name": a.trader_name,
            "trader_email": maile.get(a.trader_id), "product_key": a.product_key,
            "phase": a.phase, "status": a.status, "source": a.source,
            "initial_balance": a.initial_balance, "balance": a.balance,
            "equity": a.equity, "open_pnl": a.open_pnl,
            "profit_pct": m.get("profit_pct"),
            **dd,
            "daily_loss_used_pct": m.get("daily_loss_used_pct"),
            "overall_dd_used_pct": m.get("overall_dd_used_pct"),
            "target_equity": m.get("target_equity"),
            "daily_floor": m.get("daily_floor"),
            "overall_floor": m.get("overall_floor"),
            "profit_target_pct": m.get("profit_target_pct"),
            "max_daily_loss_pct": a.max_daily_loss_pct,
            "max_overall_loss_pct": a.max_overall_loss_pct,
            "drawdown_type": a.drawdown_type,
            "min_trading_days": a.min_trading_days,
            "trading_days": a.trading_days_count,
            **st,
            "payouts_count": len(wyp),
            "payouts_total": round(sum(p.trader_share or 0 for p in wyp), 2),
            "breaches_count": len(breache.get(a.id, [])),
            "breach_reason": a.breach_reason,
            "bot_enabled": bool(a.bot_enabled), "mt5_backed": bool(a.mt5_backed),
            "platform_server": a.platform_server,
            "created_at": _naiwna(a.created_at),
            "started_at": _naiwna(getattr(a, "started_at", None)),
        })
    _wiersze(ws, wiersze_kont, {
        "initial_balance": PIENIADZ, "balance": PIENIADZ, "equity": PIENIADZ,
        "open_pnl": PIENIADZ, "max_drawdown": PIENIADZ, "peak_equity": PIENIADZ,
        "target_equity": PIENIADZ, "daily_floor": PIENIADZ,
        "overall_floor": PIENIADZ, "gross_profit": PIENIADZ,
        "gross_loss": PIENIADZ, "net_pnl": PIENIADZ, "expectancy": PIENIADZ,
        "avg_win": PIENIADZ, "avg_loss": PIENIADZ, "largest_win": PIENIADZ,
        "largest_loss": PIENIADZ, "payouts_total": PIENIADZ,
        "profit_pct": PROCENT, "max_drawdown_pct": PROCENT,
        "win_rate_pct": PROCENT, "daily_loss_used_pct": PROCENT,
        "overall_dd_used_pct": PROCENT,
        "max_drawdown_at": DATA, "first_trade_at": DATA, "last_trade_at": DATA,
        "created_at": DATA, "started_at": DATA})

    # --- Trades -------------------------------------------------------------
    ws = wb.create_sheet("Trades")
    _naglowek(ws, ["account_login", "trader_name", "trade_id", "symbol", "side",
                   "lots", "open_price", "close_price", "pnl",
                   "pnl_pct_of_initial", "opened_at", "closed_at",
                   "duration_min", "status", "source"],
              {"opened_at": 18, "closed_at": 18, "trader_name": 22})
    wiersze_trejdow = []
    for a in konta:
        for t in trejdy_konta.get(a.id, []):
            otw, zam = _naiwna(t.opened_at), _naiwna(t.closed_at)
            wiersze_trejdow.append({
                "account_login": a.login, "trader_name": a.trader_name,
                "trade_id": t.id, "symbol": t.symbol, "side": t.side,
                "lots": t.lots, "open_price": t.open_price,
                "close_price": t.close_price, "pnl": round(t.pnl or 0, 2),
                "pnl_pct_of_initial": (round((t.pnl or 0) / a.initial_balance * 100, 4)
                                       if a.initial_balance else None),
                "opened_at": otw, "closed_at": zam,
                "duration_min": (round((zam - otw).total_seconds() / 60, 1)
                                 if otw and zam else None),
                "status": t.status, "source": t.source})
    _wiersze(ws, wiersze_trejdow, {
        "pnl": PIENIADZ, "open_price": '#,##0.00000', "close_price": '#,##0.00000',
        "pnl_pct_of_initial": '0.0000', "opened_at": DATA, "closed_at": DATA})

    # --- Daily --------------------------------------------------------------
    ws = wb.create_sheet("Daily")
    _naglowek(ws, ["account_login", "day", "trades", "wins", "losses",
                   "gross_profit", "gross_loss", "net_pnl", "lots", "symbols",
                   "equity_open", "equity_close", "balance_close",
                   "day_return_pct", "peak_equity_to_date",
                   "drawdown_from_peak_pct", "snapshots"], {"symbols": 26})
    wiersze_dni = []
    for a in konta:
        for d in dni_handlowe(trejdy_konta.get(a.id, []), krzywe.get(a.id, [])):
            wiersze_dni.append({"account_login": a.login, **d})
    _wiersze(ws, wiersze_dni, {
        "gross_profit": PIENIADZ, "gross_loss": PIENIADZ, "net_pnl": PIENIADZ,
        "equity_open": PIENIADZ, "equity_close": PIENIADZ,
        "balance_close": PIENIADZ, "peak_equity_to_date": PIENIADZ,
        "day_return_pct": '0.000', "drawdown_from_peak_pct": '0.000'})

    # --- Symbols ------------------------------------------------------------
    ws = wb.create_sheet("Symbols")
    _naglowek(ws, ["account_login", "symbol", "trades", "wins", "losses",
                   "win_rate_pct", "net_pnl", "lots", "avg_pnl"])
    wiersze_symboli = []
    for a in konta:
        po_symbolu: dict[str, list[Trade]] = defaultdict(list)
        for t in trejdy_konta.get(a.id, []):
            po_symbolu[t.symbol].append(t)
        for symbol, lista in sorted(po_symbolu.items()):
            s = statystyki_trejdow(lista)
            wiersze_symboli.append({
                "account_login": a.login, "symbol": symbol,
                "trades": s["trades_total"], "wins": s["wins"],
                "losses": s["losses"], "win_rate_pct": s["win_rate_pct"],
                "net_pnl": s["net_pnl"], "lots": s["total_lots"],
                "avg_pnl": s["expectancy"]})
    _wiersze(ws, wiersze_symboli, {"net_pnl": PIENIADZ, "avg_pnl": PIENIADZ,
                                   "win_rate_pct": PROCENT})

    # --- Equity -------------------------------------------------------------
    ws = wb.create_sheet("Equity")
    _naglowek(ws, ["account_login", "ts", "day_key", "balance", "equity",
                   "open_pnl"], {"ts": 18})
    wiersze_equity = []
    for a in konta:
        for s in krzywe.get(a.id, []):
            wiersze_equity.append({
                "account_login": a.login, "ts": _naiwna(s.ts),
                "day_key": s.day_key, "balance": round(s.balance or 0, 2),
                "equity": round(s.equity or 0, 2),
                "open_pnl": round(s.open_pnl or 0, 2)})
    _wiersze(ws, wiersze_equity, {"balance": PIENIADZ, "equity": PIENIADZ,
                                  "open_pnl": PIENIADZ, "ts": DATA})

    # --- Payouts ------------------------------------------------------------
    ws = wb.create_sheet("Payouts")
    _naglowek(ws, ["account_login", "trader_name", "ts", "profit_amount",
                   "trader_share", "paid", "method", "balance_reset",
                   "cert_token", "note"],
              {"cert_token": 26, "ts": 18, "note": 34})
    wiersze_wyplat = []
    for a in konta:
        for p in sorted(wyplaty.get(a.id, []),
                        key=lambda x: _naiwna(x.ts) or datetime.min):
            wiersze_wyplat.append({
                "account_login": a.login, "trader_name": a.trader_name,
                "ts": _naiwna(p.ts), "profit_amount": p.profit_amount,
                "trader_share": p.trader_share, "paid": bool(p.paid),
                "method": p.method, "balance_reset": p.balance_reset,
                "cert_token": p.cert_token, "note": p.note})
    _wiersze(ws, wiersze_wyplat, {"profit_amount": PIENIADZ,
                                  "trader_share": PIENIADZ,
                                  "balance_reset": PIENIADZ, "ts": DATA})

    # --- Breaches -----------------------------------------------------------
    ws = wb.create_sheet("Breaches")
    _naglowek(ws, ["account_login", "trader_name", "ts", "type", "detail",
                   "equity_at_breach"], {"detail": 44, "ts": 18})
    wiersze_breachy = []
    for a in konta:
        for b in sorted(breache.get(a.id, []),
                        key=lambda x: _naiwna(x.ts) or datetime.min):
            wiersze_breachy.append({
                "account_login": a.login, "trader_name": a.trader_name,
                "ts": _naiwna(b.ts), "type": b.type, "detail": b.detail,
                "equity_at_breach": getattr(b, "equity_at_breach", None)})
    _wiersze(ws, wiersze_breachy, {"equity_at_breach": PIENIADZ, "ts": DATA})

    bufor = io.BytesIO()
    wb.save(bufor)
    return bufor.getvalue()


def nazwa_pliku(now: datetime | None = None) -> str:
    dzis = (now or datetime.now(timezone.utc)).strftime("%Y-%m-%d")
    return f"account-statements-{dzis}.xlsx"
